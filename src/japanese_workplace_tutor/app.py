"""Streamlit application with persistent local demo authentication."""

import re
from collections.abc import Callable, Sequence
from concurrent.futures import Future, ThreadPoolExecutor
from datetime import datetime, timezone
from html import escape
from pathlib import Path

import streamlit as st
from sqlalchemy import Engine

from .auth import AuthenticationError, AuthenticationService, RegistrationError
from .database import check_database, create_database_engine
from .generation import (
    Comprehensibility,
    GenerationError,
    GeneratedQuiz,
    LessonGenerationService,
    Recovery,
    ScenarioMode,
    SlideInput,
    detect_scenario_mode,
    evaluate_delivery,
    evaluate_graduation,
)
from .core.ports import ModelTransport
from .core.priority import merge_blank_targets, suggest_personal_focus
from .core.learning_path import route_scenario_matches
from .llm.mock import MockTransport, RecordingTransport
from .prompts import lessons as lesson_prompts, presentation as presentation_prompts, roleplay as roleplay_prompts
from .review import REASONS, ReviewService, review_targets
from .llm.openai_transport import AzureOpenAITransport, OpenAICompatibleTransport
from .personal_terms import PersonalTermService, TermValidationError
from .situation_words import SituationWordService
from .domain_words import DomainWordService
from .today import TodayService, load_situation_steps, tag_today_draft
from .lesson import (
    FixtureQuestion,
    ActiveLesson,
    ActiveLessonDraft,
    ActiveReview,
    AnswerConfidence,
    FIXTURE_LESSON,
    LessonExplanationPoint,
    LessonLineExplanation,
    LessonService,
    LessonStateError,
    ProgressRecord,
    QuestionForm,
    display_options,
)
from .profile import (
    JapaneseLevel,
    ProfileRecord,
    ProfileService,
    ProfileValidationError,
    ROLE_TASK_SUGGESTIONS,
)
from .settings import Settings, get_settings
from .core.language import LanguageToolkit
from .language.sudachi_toolkit import SudachiToolkit
from .core.progress import ItemStatus, item_status
from .progress_view import build_progress_history, build_progress_view, load_milestones

PageRenderer = Callable[[], None]

JAPANESE_STYLE = """
<style>
.jp-line { font-size: 1.35rem; line-height: 2.2; margin: 0.15rem 0 0.4rem 0;
           font-family: "Yu Gothic", "Meiryo", "Noto Sans JP", "Hiragino Kaku Gothic ProN",
           sans-serif; }
.jp-line rt { font-size: 0.5em; opacity: 0.75; font-weight: 400; }
.jp-speaker { font-weight: 600; opacity: 0.7; margin-right: 0.4rem; }
</style>
"""


def _furigana_html(japanese_text: str, points: Sequence[LessonExplanationPoint]) -> str:
    """Wrap known kanji expressions in ruby tags, escaping all model text first."""

    rubies = sorted(
        (
            (point.expression, point.reading)
            for point in points
            if point.expression and point.reading
            and any("\u3400" <= character <= "\u9fff" for character in point.expression)
        ),
        key=lambda pair: len(pair[0]),
        reverse=True,
    )
    pieces: list[str] = []
    index = 0
    while index < len(japanese_text):
        for expression, reading in rubies:
            if japanese_text.startswith(expression, index):
                pieces.append(
                    f"<ruby>{escape(expression)}<rt>{escape(reading)}</rt></ruby>"
                )
                index += len(expression)
                break
        else:
            pieces.append(escape(japanese_text[index]))
            index += 1
    return "".join(pieces)


def _render_japanese_line(
    japanese_text: str,
    points: Sequence[LessonExplanationPoint],
    show_furigana: bool,
) -> None:
    speaker, separator, spoken = japanese_text.partition(": ")
    if not separator:
        speaker, spoken = "", japanese_text
    body = (
        _furigana_html(spoken, points) if show_furigana else escape(spoken)
    )
    prefix = f'<span class="jp-speaker">{escape(speaker)}</span>' if speaker else ""
    st.markdown(
        f'<div class="jp-line" lang="ja">{prefix}{body}</div>',
        unsafe_allow_html=True,
    )


def _line_points(line: LessonLineExplanation) -> tuple[LessonExplanationPoint, ...]:
    return (*line.kanji, *line.vocabulary, *line.grammar)


def _confidence_input(key: str, disabled: bool) -> AnswerConfidence:
    choice = st.radio(
        "How sure are you?",
        ("I am sure", "I guessed"),
        index=0,
        key=key,
        horizontal=True,
        disabled=disabled,
    )
    return (
        AnswerConfidence.GUESSED if choice == "I guessed" else AnswerConfidence.SURE
    )


def _glossary(
    line_explanations: Sequence[LessonLineExplanation],
) -> list[tuple[str, LessonExplanationPoint, list[int]]]:
    """Collapse repeated expressions into one entry listing the lines they appear on."""

    entries: dict[tuple[str, str], tuple[str, LessonExplanationPoint, list[int]]] = {}
    for number, line in enumerate(line_explanations, start=1):
        for label, points in (
            ("Kanji", line.kanji),
            ("Vocabulary", line.vocabulary),
            ("Grammar", line.grammar),
        ):
            for point in points:
                key = (label, point.expression)
                if key not in entries:
                    entries[key] = (label, point, [])
                entries[key][2].append(number)
    return list(entries.values())


PUBLIC_CONTENT_NOTE = (
    "Do not enter or upload work data: no customer names, project details, internal "
    "numbers or anything confidential. Text you submit is sent to an external AI service."
)


def _public_content_confirmed(key: str) -> bool:
    """Decision 2: user-supplied text needs an explicit public-content confirmation."""

    st.warning(PUBLIC_CONTENT_NOTE)
    return st.checkbox(
        "I confirm this content is public and non-confidential", key=key, value=False
    )


def _review_versions(
    generation_service: "LessonGenerationService | None", active: ActiveLesson
) -> dict[str, object]:
    return {
        "prompts": {
            "lessons": lesson_prompts.PROMPT_VERSION,
            "roleplay": roleplay_prompts.PROMPT_VERSION,
            "presentation": presentation_prompts.PROMPT_VERSION,
        },
        "toolkit": generation_service.toolkit.version if generation_service else None,
        "lesson_provider": active.provider_name,
        "quiz_provider": active.quiz_provider_name,
        "policy": LessonService.POLICY_VERSION,
    }


def _render_review_panel(
    review_service: ReviewService,
    user_id: int,
    active: ActiveLesson,
    versions: dict[str, object],
) -> None:
    """Reviewer mode: mark each line, item and question right or wrong (decision 9)."""

    with st.expander("Reviewer: check this lesson's accuracy", expanded=False):
        st.caption(
            "Mark only what you checked. Verdicts and the reviewed text are saved for the "
            "accuracy report."
        )
        choices: list[tuple[object, str, str | None, str]] = []
        for target in review_targets(active):
            key = f"review_{active.lesson_session_id}_{target.kind}_{target.target_id}"
            st.markdown(f"**{target.kind.title()}** · {target.label}")
            verdict = st.radio(
                "Verdict",
                ("Not reviewed", "Right", "Wrong"),
                horizontal=True,
                key=key,
                label_visibility="collapsed",
            )
            reason = note = None
            if verdict == "Wrong":
                reason = st.selectbox(
                    "What is wrong?", tuple(REASONS), format_func=REASONS.__getitem__,
                    key=f"{key}_reason",
                )
                note = st.text_input("Note (optional)", key=f"{key}_note")
            if verdict != "Not reviewed":
                choices.append((target, "ok" if verdict == "Right" else "wrong", reason, note or ""))
        if st.button("Save review", key=f"save_review_{active.lesson_session_id}"):
            for target, verdict, reason, note in choices:
                review_service.record(
                    user_id, active.lesson_session_id, target, verdict,
                    reason=reason, note=note, versions=versions,
                )
            st.success(f"Saved {len(choices)} verdicts.")


def render_accuracy(review_service: ReviewService) -> None:
    st.title("Accuracy")
    report = review_service.report()
    if not report.reviewed:
        st.info("No reviews yet. Open a generated lesson and use the reviewer panel.")
        return
    left, middle, right = st.columns(3)
    left.metric("Reviewed", report.reviewed)
    middle.metric("Marked wrong", report.wrong)
    right.metric("Error rate", f"{report.error_rate:.1%}")
    st.subheader("By kind")
    st.table(
        [{"kind": kind, "wrong": wrong, "reviewed": total, "error rate": f"{wrong / total:.1%}"}
         for kind, (wrong, total) in sorted(report.by_kind.items())]
    )
    if report.by_key_source:
        st.subheader("Questions by answer-key source")
        st.caption("Derived keys are graded; unverified keys are practice only.")
        st.table(
            [{"key source": source, "wrong": wrong, "reviewed": total,
              "error rate": f"{wrong / total:.1%}"}
             for source, (wrong, total) in sorted(report.by_key_source.items())]
        )
    if report.by_reason:
        st.subheader("Why it was wrong")
        st.table([{"reason": REASONS.get(reason, reason), "count": count}
                  for reason, count in sorted(report.by_reason.items(), key=lambda kv: -kv[1])])
    st.subheader("Everything marked wrong")
    for row in report.wrong_rows:
        st.markdown(f"**{row['kind']}** `{row['target']}` · {REASONS.get(row['reason'] or 'other')}")
        if row["note"]:
            st.caption(row["note"])
        st.json({"snapshot": row["snapshot"], "versions": row["versions"]}, expanded=False)


def _question_caption(question: FixtureQuestion) -> str:
    form = question.form.value.replace("_", " ").title()
    if question.graded:
        return form
    return f"{form} · Practice: the answer could not be checked automatically, so it does not count toward progress"


@st.cache_resource
def _quiz_executor() -> ThreadPoolExecutor:
    return ThreadPoolExecutor(max_workers=2, thread_name_prefix="quiz-generation")


@st.cache_resource
def _generation_service(
    base_url: str,
    _api_key: str,
    primary_model: str,
    fallback_model: str,
    timeout_seconds: float,
    primary_timeout_seconds: float,
    hedge_after_seconds: float | None,
    parallel_explanations: bool,
    explanation_workers: int,
    quiz_primary_model: str,
    quiz_fallback_model: str,
    llm_mode: str = "live",
    recording_path: str = "",
    model_provider: str = "openai-compatible",
    azure_api_version: str | None = None,
    _apim_subscription_key: str | None = None,
    apim_subscription_key_header: str = "Ocp-Apim-Subscription-Key",
) -> LessonGenerationService:
    transport: ModelTransport
    if model_provider == "azure":
        transport = AzureOpenAITransport(
            base_url, _api_key, timeout_seconds, azure_api_version,
            apim_subscription_key=_apim_subscription_key,
            apim_subscription_key_header=apim_subscription_key_header,
        )
    else:
        transport = OpenAICompatibleTransport(base_url, _api_key, timeout_seconds)
    if llm_mode == "record":
        transport = RecordingTransport(transport, Path(recording_path))
    return LessonGenerationService(
        transport,
        primary_model,
        fallback_model,
        primary_timeout_seconds=primary_timeout_seconds,
        hedge_after_seconds=hedge_after_seconds,
        parallel_explanations=parallel_explanations,
        explanation_workers=explanation_workers,
        quiz_primary_model=quiz_primary_model,
        quiz_fallback_model=quiz_fallback_model,
    )


@st.cache_resource
def _mock_generation_service(fixture_dir: str) -> LessonGenerationService:
    """Offline service: real prompts, validators and fallback over replayed fixtures."""

    return LessonGenerationService(
        MockTransport(Path(fixture_dir)), "mock-primary", "mock-fallback"
    )


def _start_quiz_generation(
    generation_service: LessonGenerationService,
    draft: ActiveLessonDraft,
    learning_history: tuple[ProgressRecord, ...],
) -> Future[GeneratedQuiz]:
    return _quiz_executor().submit(
        generation_service.generate_quiz,
        draft.content,
        learning_history,
    )


def _clear_active_review() -> None:
    for key in list(st.session_state):
        if key in {"active_due_review", "due_review_results"} or key.startswith(
            "due_review_"
        ):
            st.session_state.pop(key, None)


def _continue_learning() -> None:
    _clear_active_review()
    st.session_state.pending_navigation = "Learn"


def _render_due_review(
    service: LessonService, user_id: int, active: ActiveReview
) -> None:
    st.subheader("Review due items")
    st.caption(f"{len(active.items)} item{'s' if len(active.items) != 1 else ''} in this review")
    results = st.session_state.setdefault("due_review_results", {})

    for number, review_item in enumerate(active.items, start=1):
        question = review_item.question
        st.markdown(
            f"**{number}. {review_item.item.expression} · "
            f"{review_item.item.category.value}**"
        )
        st.write(question.prompt)
        st.caption(question.form.value.replace("_", " ").title())
        answer_key = (
            f"due_review_{active.review_session_id}_{question.question_id}"
        )
        selected_answer = st.radio(
            f"Answer for review item {number}",
            display_options(question),
            index=None,
            key=answer_key,
            label_visibility="collapsed",
            disabled=question.question_id in results,
        )
        confidence = _confidence_input(
            f"due_review_confidence_{active.review_session_id}_{question.question_id}",
            question.question_id in results,
        )
        if question.question_id in results:
            result = results[question.question_id]
            if result.is_correct:
                st.success(f"Correct. Outcome: {result.outcome.value.title()}.")
            else:
                correct_answer = question.options[question.correct_option_index]
                st.error(f"Not quite. The correct answer is {correct_answer}.")
            st.write(question.explanation)
            st.caption(
                "Next review: "
                + result.next_review_at.isoformat(sep=" ", timespec="minutes")
            )
        elif st.button(
            "Submit answer",
            key=(
                f"submit_due_review_{active.review_session_id}_"
                f"{question.question_id}"
            ),
        ):
            if selected_answer is None:
                st.warning("Choose an answer before submitting.")
            else:
                results[question.question_id] = service.submit_review_answer(
                    user_id,
                    active.review_session_id,
                    question.question_id,
                    question.options.index(selected_answer),
                    confidence=confidence,
                    active_review=active,
                )
                st.rerun()

    if active.items and len(results) == len(active.items):
        if st.button("Complete review", type="primary", key="complete_due_review"):
            st.session_state.last_review_summary = {
                "correct": sum(result.is_correct for result in results.values()),
                "total": len(active.items),
                "next_review_at": service.get_next_review_at(user_id),
            }
            _clear_active_review()
            if st.session_state.get("today_route"):
                _continue_learning()
            st.rerun()

    st.button("Start a lesson", key="continue_learning", on_click=_continue_learning)
    if st.button("Skip review", key="skip_due_review"):
        _clear_active_review()
        if st.session_state.get("today_route"):
            _continue_learning()
        st.rerun()


def render_home(
    profile: ProfileRecord | None = None,
    service: LessonService | None = None,
    user_id: int | None = None,
    today_service: TodayService | None = None,
) -> None:
    st.title("Home")
    username = st.session_state.get("authenticated_username")
    st.write(f"Welcome, {username}.")
    if profile is None:
        st.info("Complete your learner profile to personalize your learning.")
        return
    st.subheader(f"{profile.role} · {profile.declared_level.value}")
    st.write("Focus tasks: " + ", ".join(profile.tasks))
    if profile.tools_domain:
        st.caption(f"Tools and domain: {profile.tools_domain}")
    if service is None or user_id is None:
        st.info("A deterministic workplace status-update lesson is ready on Learn.")
        return

    active_review = st.session_state.get("active_due_review")
    if isinstance(active_review, ActiveReview):
        _render_due_review(service, user_id, active_review)
        return

    _render_this_week(service, user_id)

    plan = today_service.build_today_plan(user_id) if today_service else None
    if plan is not None:
        st.subheader("Today")
        if plan.next_step is not None:
            st.write(plan.next_step.title)
            st.caption(plan.reason)
            st.caption(f"Route steps completed: {plan.completed_steps} of {plan.total_steps}")
            if plan.review_count:
                st.caption(f"First, review up to {plan.review_count} due items. You can skip them.")
            if st.button("Start today", type="primary", key="start_today"):
                st.session_state.today_route = plan.next_step.topic_id
                st.session_state.today_scenario = plan.next_step.scenario
                st.session_state.lesson_scenario = plan.next_step.scenario
                if plan.review_count:
                    st.session_state.active_due_review = service.start_due_review(user_id)
                    st.session_state.due_review_results = {}
                else:
                    _continue_learning()
                st.rerun()
        else:
            st.write(plan.reason)
        if plan.local_alternative is not None and plan.local_alternative != plan.next_step:
            st.caption(f"Practice offline: {plan.local_alternative.title}")
            if st.button("Practice offline", key="start_local_practice"):
                st.session_state.today_route = plan.local_alternative.topic_id
                st.session_state.today_scenario = ""
                st.session_state.lesson_scenario = ""
                _continue_learning()
                st.rerun()

    summary = st.session_state.get("last_review_summary")
    if isinstance(summary, dict):
        st.success(
            f"Review complete: {summary['correct']} of {summary['total']} correct."
        )
        next_review_at = summary.get("next_review_at")
        if next_review_at is not None:
            st.caption(
                "Earliest review: "
                + next_review_at.isoformat(sep=" ", timespec="minutes")
            )

    due_count = service.get_due_count(user_id)
    if due_count:
        st.subheader(f"{due_count} item{'s' if due_count != 1 else ''} due")
        st.write("A short review is ready. You can skip it and keep learning.")
        if st.button("Review due items", type="primary", key="start_due_review"):
            st.session_state.pop("last_review_summary", None)
            st.session_state.active_due_review = service.start_due_review(user_id)
            st.session_state.due_review_results = {}
            st.rerun()
        st.button(
            "Start a lesson", key="continue_learning", on_click=_continue_learning
        )
    else:
        st.write("No reviews are due. Continue with the workplace status-update lesson.")
        st.button(
            "Continue learning",
            type="primary",
            key="continue_learning",
            on_click=_continue_learning,
        )


def _clear_active_lesson() -> None:
    future = st.session_state.get("quiz_generation_future")
    if isinstance(future, Future):
        future.cancel()
    for key in list(st.session_state):
        if key in {
            "active_fixture_lesson",
            "active_lesson",
            "active_lesson_draft",
            "lesson_answer_results",
            "lesson_retry_results",
            "quiz_visible",
            "lesson_fallback_notice",
            "quiz_generation_future",
            "quiz_generation_error",
        } or key.startswith(("lesson_answer_", "lesson_retry_")):
            st.session_state.pop(key, None)
    _clear_roleplay()


def _clear_roleplay() -> None:
    for key in list(st.session_state):
        if key.startswith("roleplay_"):
            st.session_state.pop(key, None)


def _roleplay_reply_input(slot: str, options: tuple[str, ...], guided: bool) -> str:
    """Guided learners pick a line; higher levels type free Japanese."""

    if guided:
        choice = st.radio(
            "Your reply",
            options,
            index=None,
            key=f"roleplay_{slot}_choice",
            label_visibility="collapsed",
        )
        return choice or ""
    return st.text_area(
        "Your reply in Japanese",
        key=f"roleplay_{slot}_text",
        label_visibility="collapsed",
        height=68,
    ).strip()


def render_roleplay(
    generation_service: LessonGenerationService | None,
    profile: ProfileRecord | None,
) -> None:
    """Focused takeover screen for the situation roleplay stage."""

    challenge = st.session_state.get("roleplay_challenge")
    stage = st.session_state.get("roleplay_stage")
    st.title("Conversation practice (beta)")
    if st.button("Leave conversation practice", key="leave_roleplay"):
        _clear_roleplay()
        st.rerun()
    if challenge is None or profile is None:
        st.info("This conversation practice is no longer available.")
        return

    # Brief: the roadmap, shown before practice so the finish line is visible.
    st.subheader(challenge.goal)
    st.caption(
        f"You: {challenge.learner_role}  ·  Them: {challenge.counterpart_role}  ·  "
        f"{challenge.setting}"
    )
    st.markdown("**Aim to:**")
    for criterion in challenge.success_criteria:
        st.markdown(f"- {criterion}")
    with st.expander("Toolkit: phrases you can use"):
        for phrase in challenge.toolkit:
            label = "recovery" if phrase.kind == "recovery" else "target"
            st.markdown(
                f"**{phrase.japanese}** ({phrase.reading}) — {phrase.meaning}  ·  _{label}_"
            )

    if stage == "brief":
        if st.button("Start roleplay", type="primary", key="begin_roleplay"):
            st.session_state.roleplay_stage = "respond"
            st.rerun()
        return

    level = profile.estimated_working_level or profile.declared_level
    guided = level in (
        JapaneseLevel.COMPLETE_BEGINNER,
        JapaneseLevel.N5,
        JapaneseLevel.N4,
        JapaneseLevel.UNSURE,
    )

    st.divider()
    st.markdown(f"**{challenge.counterpart_role} says:**")
    st.markdown(f"> {challenge.opening_line}")
    opening_reply = _roleplay_reply_input("open", challenge.opening_options, guided)

    st.markdown(f"**Then, unexpectedly ({challenge.curveball_type}):**")
    st.markdown(f"> {challenge.curveball_line}")
    curveball_reply = _roleplay_reply_input("curve", challenge.curveball_options, guided)

    if stage == "respond":
        if st.button("Submit for feedback", type="primary", key="submit_roleplay"):
            if not opening_reply or not curveball_reply:
                st.warning("Answer both lines before submitting.")
            elif generation_service is None:
                st.error("Feedback is unavailable until model configuration is complete.")
            else:
                with st.spinner("Evaluating your conversation..."):
                    try:
                        judgement = generation_service.judge_roleplay(
                            challenge, opening_reply, curveball_reply
                        )
                    except GenerationError as error:
                        st.error(str(error))
                    else:
                        st.session_state.roleplay_verdict = judgement.verdict
                        st.session_state.roleplay_provider = judgement.provider_name
                        st.session_state.roleplay_stage = "verdict"
                        st.rerun()
        return

    # Verdict
    verdict = st.session_state.get("roleplay_verdict")
    if verdict is None:
        return
    result = evaluate_graduation(verdict)
    st.divider()
    st.subheader("Feedback")

    def _row(label: str, passed: bool, detail: str) -> None:
        st.markdown(f"{'✅' if passed else '❌'} **{label}** — {detail}")

    _row("Goal achieved", verdict.goal_achieved, verdict.goal_evidence)
    _row(
        "Comprehensible",
        verdict.comprehensible != Comprehensibility.NO,
        f"{verdict.comprehensible.value}: {verdict.comprehensible_evidence}",
    )
    _row("Right register", verdict.register_appropriate, verdict.register_evidence)
    _row(
        "Recovered from the surprise",
        verdict.recovered_from_curveball != Recovery.NO,
        f"{verdict.recovered_from_curveball.value}: {verdict.recovery_evidence}",
    )
    provider = st.session_state.get("roleplay_provider")
    st.caption(
        "AI feedback (beta): it can be wrong, and it is not saved to your progress."
        + (f" Evaluated with {provider}." if provider else "")
    )
    if result.failed_criteria:
        st.info("The AI suggests focusing on: " + ", ".join(result.failed_criteria) + ".")
    # Decision 6: no stored pass/fail. The learner decides when they are ready.
    if st.checkbox("I feel ready to do this for real", key="roleplay_self_check"):
        st.markdown(
            f"**Now try it for real.** This week, actually {challenge.goal[0].lower()}"
            f"{challenge.goal[1:]}. Then come back and practise another situation."
        )
    if st.button("Try again", key="retry_roleplay"):
        for key in ("roleplay_verdict", "roleplay_provider", "roleplay_self_check"):
            st.session_state.pop(key, None)
        st.session_state.roleplay_stage = "respond"
        st.rerun()


def _render_lesson_study(lesson, line_explanations) -> None:
    target_expressions = {item.expression for item in lesson.items}
    points_by_line = {
        line.japanese_text: _line_points(line) for line in line_explanations
    }
    furigana_column, english_column = st.columns(2)
    show_furigana = furigana_column.toggle("Furigana", value=True, key="show_furigana")
    show_english = english_column.toggle("Show English", value=False, key="show_english")

    conversation_tab, breakdown_tab, glossary_tab = st.tabs(
        ("Conversation", "Line breakdown", "Glossary")
    )
    with conversation_tab:
        for number, dialogue_line in enumerate(
            (line for line in lesson.passage.splitlines() if line.strip()), start=1
        ):
            _render_japanese_line(
                dialogue_line.strip(),
                points_by_line.get(dialogue_line.strip(), ()),
                show_furigana,
            )
            if show_english and number <= len(line_explanations):
                st.caption(line_explanations[number - 1].english_meaning)

    with breakdown_tab:
        if not line_explanations:
            st.caption("This lesson has no line-by-line explanation.")
        for number, line in enumerate(line_explanations, start=1):
            introduced = _line_points(line)
            summary = ", ".join(point.expression for point in introduced) or "review"
            with st.expander(f"Line {number} — {summary}"):
                _render_japanese_line(line.japanese_text, introduced, show_furigana)
                st.write(line.english_meaning)
                for label, points in (
                    ("Kanji", line.kanji),
                    ("Vocabulary", line.vocabulary),
                    ("Grammar", line.grammar),
                ):
                    for point in points:
                        badge = " ⭐ practice target" if point.expression in target_expressions else ""
                        st.markdown(
                            f"**{label} · {point.expression}** ({point.reading}){badge}"
                        )
                        st.write(f"{point.meaning}. {point.explanation}")

    with glossary_tab:
        glossary = _glossary(line_explanations)
        if not glossary:
            st.caption("This lesson has no glossary entries.")
        for label, point, line_numbers in glossary:
            lines_text = ", ".join(str(number) for number in line_numbers)
            badge = " ⭐" if point.expression in target_expressions else ""
            st.markdown(
                f"**{point.expression}**{badge} ({point.reading}) · {label} · "
                f"line {lines_text}"
            )
            st.write(f"{point.meaning}. {point.explanation}")


def _render_today_practice(today_service: TodayService, step) -> None:
    st.subheader(step.title)
    if step.track == "grammar":
        question = today_service.grammar_question(step)
        if question is None:
            st.warning("This grammar practice is not available yet.")
            return
        meaning, options = question
        answer = st.radio(
            f"Which Japanese form means '{meaning}'?", options,
            index=None, key="today_grammar_answer",
        )
    else:
        st.write("Put the chunks in order to build the workplace phrase.")
        st.caption("Chunks: " + " / ".join(reversed(step.chunks)))
        selections = tuple(st.selectbox(
            f"Position {index + 1}", tuple(reversed(step.chunks)), index=None,
            key=f"today_chunk_{index}",
        ) for index in range(len(step.chunks)))
        answer = (tuple(step.chunks.index(chunk) for chunk in selections)
                  if all(chunk is not None for chunk in selections) else None)
    if st.button("Check answer", key="submit_today_practice"):
        if answer is None:
            st.warning("Choose an answer first.")
        elif today_service.submit_practice(
            int(st.session_state.authenticated_user_id), step.topic_id, answer,
        ):
            st.session_state.pop("today_route", None)
            st.session_state.pop("today_scenario", None)
            st.session_state.today_practice_complete = True
            st.rerun()
        else:
            st.error("Try another answer.")


def render_learn(
    service: LessonService | None = None,
    user_id: int | None = None,
    profile: ProfileRecord | None = None,
    generation_service: LessonGenerationService | None = None,
    review_service: ReviewService | None = None,
    term_service: PersonalTermService | None = None,
    situation_service: SituationWordService | None = None,
    domain_service: DomainWordService | None = None,
    llm_mode: str = "live",
    today_service: TodayService | None = None,
) -> None:
    st.title("Learn")
    if service is None or user_id is None:
        st.info("Sign in to start a lesson.")
        return

    if st.session_state.pop("today_practice_complete", False):
        st.success("Practice complete. Your next Today step is ready on Home.")

    route_id = st.session_state.get("today_route")
    route_step = next((step for step in load_situation_steps()
                       if step.topic_id == route_id), None)
    if (route_step is not None and route_step.track in {"grammar", "sentence-order"}
            and today_service is not None):
        _render_today_practice(today_service, route_step)
        return

    if llm_mode == "mock" and route_id:
        st.info(
            "Offline sample lessons do not complete a Today generated lesson route. "
            "Configure a model to advance this route."
        )

    if "last_lesson_completion" in st.session_state:
        st.success("Lesson completed. Your compact progress and review schedule were saved.")
        st.session_state.pop("last_lesson_completion", None)

    if st.session_state.get("roleplay_stage"):
        render_roleplay(generation_service, profile)
        return

    active = st.session_state.get("active_lesson")
    draft = st.session_state.get("active_lesson_draft")
    if not isinstance(active, ActiveLesson) and not isinstance(draft, ActiveLessonDraft):
        st.subheader("Create a scenario lesson")
        scenario = st.text_area(
            "Workplace situation, goal, or Japanese text",
            max_chars=4000,
            key="lesson_scenario",
            placeholder="For example: I need to clarify a requirement with my manager.",
        )
        detected_mode = detect_scenario_mode(scenario) if scenario.strip() else ScenarioMode.GENERATE
        mode_options = [mode.value for mode in ScenarioMode]
        selected_mode = st.radio(
            "Lesson mode",
            mode_options,
            index=mode_options.index(detected_mode.value),
            key=f"scenario_mode_{detected_mode.name.lower()}",
            horizontal=True,
        )
        st.caption(f"Detected/default mode: {detected_mode.value}. Confirm or change it above.")
        completions = service.get_completions(user_id)
        include_personal = (
            st.checkbox(
                "Focus this surprise lesson on one of my saved words",
                value=suggest_personal_focus(len(completions)),
                key=f"personal_focus_{len(completions)}",
            ) if term_service is not None else False
        )
        generation_disabled = generation_service is None or profile is None
        if generation_disabled:
            st.info("Scenario generation is unavailable until model configuration is complete.")
        if scenario.strip() and not _public_content_confirmed("scenario_public_ok"):
            generation_disabled = True
        if st.button(
            "Generate lesson",
            type="primary",
            key="generate_scenario_lesson",
            disabled=generation_disabled,
        ) and generation_service is not None and profile is not None:
            try:
                learning_history = service.get_progress(user_id)
                now = datetime.now(timezone.utc).replace(tzinfo=None)
                selected_targets = ()
                if ScenarioMode(selected_mode) is ScenarioMode.GENERATE:
                    if (domain_service is not None and route_scenario_matches(
                            route_step, st.session_state.get("today_scenario"), scenario)):
                        selected_targets = domain_service.route_targets(
                            route_step.words, profile, tuple(learning_history), now
                        )
                    elif scenario.strip() and situation_service is not None:
                        selected_targets = situation_service.lesson_targets(
                            scenario, profile, tuple(learning_history), now
                        )
                    elif not scenario.strip():
                        personal_targets = (
                            term_service.lesson_targets(
                                user_id, profile, tuple(learning_history), now
                            ) if term_service is not None and include_personal else ()
                        )
                        domain_targets = (
                            domain_service.lesson_targets(
                                scenario, profile, tuple(learning_history), now
                            ) if domain_service is not None else ()
                        )
                        selected_targets = merge_blank_targets(
                            personal_targets, domain_targets
                        )
                generated = generation_service.generate_lesson_content(
                    scenario,
                    ScenarioMode(selected_mode),
                    profile,
                    learning_history=learning_history,
                    recent_topic_ids=tuple(
                        completion.topic_id
                        for completion in completions[:10]
                    ),
                    selected_targets=selected_targets,
                )
                generated = tag_today_draft(
                    generated,
                    (route_id if ScenarioMode(selected_mode) is ScenarioMode.GENERATE
                     and (route_step is None or route_step.track != "words"
                          or selected_targets) else None),
                    st.session_state.get("today_scenario"),
                    scenario,
                    mock_mode=llm_mode == "mock",
                )
            except GenerationError as error:
                # Always-ship fallback: never dead-end the demo. If every model and
                # repair attempt fails, load the ready-made fixture lesson so the
                # learner can still complete the full loop, with an honest notice.
                _clear_active_lesson()
                st.session_state.pop("today_route", None)
                st.session_state.pop("today_scenario", None)
                st.session_state.active_lesson = service.start_fixture_lesson(user_id)
                st.session_state.quiz_visible = False
                st.session_state.lesson_answer_results = {}
                st.session_state.lesson_retry_results = {}
                st.session_state.lesson_fallback_notice = (
                    "Live lesson generation is temporarily unavailable, so we loaded a "
                    "ready-made sample lesson so you can keep learning. Please try "
                    f"generating again in a moment. (Reason: {error})"
                )
                st.rerun()
            else:
                st.session_state.pop("today_route", None)
                st.session_state.pop("today_scenario", None)
                _clear_active_lesson()
                active_draft = service.start_generated_lesson_draft(
                    user_id, generated
                )
                st.session_state.active_lesson_draft = active_draft
                st.session_state.quiz_generation_future = _start_quiz_generation(
                    generation_service,
                    active_draft,
                    tuple(learning_history),
                )
                st.session_state.lesson_answer_results = {}
                st.session_state.lesson_retry_results = {}
                st.rerun()

        st.divider()
        st.subheader(FIXTURE_LESSON.title)
        st.write("Study five workplace Japanese targets, then answer five questions.")
        st.caption("Opening the lesson records exposure only. It does not raise mastery.")
        if st.button("Start fixture lesson", key="start_fixture_lesson"):
            _clear_active_lesson()
            st.session_state.pop("today_route", None)
            st.session_state.pop("today_scenario", None)
            st.session_state.active_lesson = service.start_fixture_lesson(
                user_id
            )
            st.session_state.quiz_visible = False
            st.session_state.lesson_answer_results = {}
            st.session_state.lesson_retry_results = {}
            st.rerun()
        return

    if isinstance(active, ActiveLesson):
        lesson = active.lesson
        line_explanations = active.line_explanations
        provider_name = active.provider_name
    else:
        lesson = draft.content
        line_explanations = draft.line_explanations
        provider_name = draft.provider_name
    fallback_notice = st.session_state.get("lesson_fallback_notice")
    if isinstance(fallback_notice, str):
        st.warning(fallback_notice)
    st.subheader(lesson.title)
    st.caption(f"Difficulty: {lesson.difficulty}")
    if provider_name is not None:
        st.caption(f"Lesson generated with {provider_name}")

    quiz_visible = isinstance(active, ActiveLesson) and st.session_state.get(
        "quiz_visible", True
    )
    if not quiz_visible:
        _render_lesson_study(lesson, line_explanations)

    if isinstance(active, ActiveLesson) and not quiz_visible:
        if st.button("Quiz", type="primary", key="go_to_fixture_quiz"):
            st.session_state.quiz_visible = True
            st.rerun()
        return

    if isinstance(draft, ActiveLessonDraft):
        future = st.session_state.get("quiz_generation_future")
        quiz_error = st.session_state.get("quiz_generation_error")
        if isinstance(quiz_error, str):
            st.error(quiz_error)
            if generation_service is not None and st.button(
                "Retry quiz preparation", key="retry_quiz_generation"
            ):
                st.session_state.pop("quiz_generation_error", None)
                st.session_state.quiz_generation_future = _start_quiz_generation(
                    generation_service,
                    draft,
                    tuple(service.get_progress(user_id)),
                )
                st.rerun()
        elif isinstance(future, Future):
            if future.done():
                st.caption("Your personalized quiz is ready.")
            else:
                st.caption("Your personalized quiz is being prepared while you study.")
            if st.button("Quiz", type="primary", key="go_to_quiz"):
                with st.spinner("Preparing your personalized quiz..."):
                    try:
                        generated_quiz = future.result()
                    except Exception:
                        st.session_state.pop("quiz_generation_future", None)
                        st.session_state.quiz_generation_error = (
                            "We could not prepare a valid quiz. Your lesson is still available."
                        )
                    else:
                        st.session_state.active_lesson = (
                            service.activate_generated_quiz(draft, generated_quiz)
                        )
                        st.session_state.quiz_visible = True
                        st.session_state.pop("active_lesson_draft", None)
                        st.session_state.pop("quiz_generation_future", None)
                    st.rerun()
        else:
            st.error("Quiz preparation was interrupted. Your lesson is still available.")
        if st.button("Leave lesson", key="leave_generated_lesson_draft"):
            _clear_active_lesson()
            st.rerun()
        return

    st.subheader("Practice")
    if active.quiz_provider_name is not None:
        st.caption(f"Quiz generated with {active.quiz_provider_name}")
    answer_results = st.session_state.setdefault("lesson_answer_results", {})
    for number, question in enumerate(lesson.questions, start=1):
        st.markdown(f"**{number}. {question.prompt}**")
        st.caption(_question_caption(question))
        answer_key = f"lesson_answer_{active.lesson_session_id}_{question.question_id}"
        answered = question.question_id in answer_results
        selected_answer = st.radio(
            f"Answer for question {number}",
            display_options(question),
            index=None,
            key=answer_key,
            label_visibility="collapsed",
            disabled=answered,
        )
        confidence = _confidence_input(
            f"lesson_confidence_{active.lesson_session_id}_{question.question_id}",
            answered,
        )
        if answered:
            if answer_results[question.question_id]:
                st.success("Correct.")
            else:
                correct_answer = question.options[question.correct_option_index]
                st.error(f"Not quite. The correct answer is {correct_answer}.")
            st.write(question.explanation)
        elif st.button(
            "Submit answer", key=f"submit_{active.lesson_session_id}_{question.question_id}"
        ):
            if selected_answer is None:
                st.warning("Choose an answer before submitting.")
            else:
                result = service.submit_answer(
                    user_id,
                    active.lesson_session_id,
                    question.question_id,
                    question.options.index(selected_answer),
                    active_lesson=active,
                    confidence=confidence,
                )
                answer_results[question.question_id] = result.is_correct
                st.rerun()

    if len(answer_results) == len(lesson.questions):
        required_retries = [
            retry
            for question, retry in zip(
                lesson.questions, active.retry_questions, strict=True
            )
            if answer_results.get(question.question_id) is False
        ]
        retry_results = st.session_state.setdefault("lesson_retry_results", {})
        if required_retries:
            st.subheader("Corrective practice")
            st.caption(
                "These alternate question forms revisit concepts missed earlier. "
                "Recovery is stored separately from the original answer."
            )
        questions_by_item = {
            question.item_id: question for question in lesson.questions
        }
        for number, retry in enumerate(required_retries, start=1):
            missed = questions_by_item.get(retry.item_id)
            if missed is not None:
                st.info(f"Before you retry: {missed.explanation}")
            st.markdown(f"**Retry {number}. {retry.prompt}**")
            st.caption(_question_caption(retry))
            retry_key = f"lesson_retry_{active.lesson_session_id}_{retry.question_id}"
            retried = retry.question_id in retry_results
            selected_retry = st.radio(
                f"Answer for retry {number}",
                display_options(retry),
                index=None,
                key=retry_key,
                label_visibility="collapsed",
                disabled=retried,
            )
            retry_confidence = _confidence_input(
                f"retry_confidence_{active.lesson_session_id}_{retry.question_id}",
                retried,
            )
            if retried:
                if retry_results[retry.question_id]:
                    st.success("Recovered. This counts as Hard, not a first-try success.")
                else:
                    correct_answer = retry.options[retry.correct_option_index]
                    st.error(f"Not quite. The correct answer is {correct_answer}.")
                st.write(retry.explanation)
            elif st.button(
                "Submit retry",
                key=f"submit_retry_{active.lesson_session_id}_{retry.question_id}",
            ):
                if selected_retry is None:
                    st.warning("Choose an answer before submitting.")
                else:
                    result = service.submit_retry_answer(
                        user_id,
                        active.lesson_session_id,
                        retry.question_id,
                        retry.options.index(selected_retry),
                        active_lesson=active,
                        confidence=retry_confidence,
                    )
                    retry_results[retry.question_id] = result.is_correct
                    st.rerun()

    required_retry_count = sum(
        answer_results.get(question.question_id) is False
        for question in lesson.questions
    )
    completed_retry_count = len(st.session_state.get("lesson_retry_results", {}))
    if (
        len(answer_results) == len(lesson.questions)
        and completed_retry_count == required_retry_count
    ):
        st.subheader("Recap")
        st.write(lesson.recap)
        if st.button("Complete lesson", type="primary", key="complete_fixture_lesson"):
            try:
                service.complete_lesson(user_id, active)
            except LessonStateError as error:
                st.error(str(error))
            else:
                _clear_active_lesson()
                st.session_state.last_lesson_completion = True
                st.rerun()

        st.divider()
        st.subheader("Practice the conversation")
        st.caption(
            "Roleplay this situation and handle one curveball, then get feedback on whether "
            "you are ready to have it for real."
        )
        if generation_service is not None and profile is not None:
            if st.button("Practice this conversation", key="start_roleplay"):
                with st.spinner("Preparing your conversation challenge..."):
                    try:
                        generated = generation_service.generate_situation_challenge(
                            lesson, profile
                        )
                    except GenerationError as error:
                        st.error(str(error))
                    else:
                        _clear_roleplay()
                        st.session_state.roleplay_challenge = generated.challenge
                        st.session_state.roleplay_provider = generated.provider_name
                        st.session_state.roleplay_stage = "brief"
                        st.rerun()
        else:
            st.info("Conversation practice needs model configuration.")

    if review_service is not None:
        _render_review_panel(
            review_service, user_id, active, _review_versions(generation_service, active)
        )

    if st.button("Leave lesson", key="leave_fixture_lesson"):
        _clear_active_lesson()
        st.rerun()


PRESENTATION_SLIDE_CAP = 6

# Display label -> deck_type key understood by DECK_TYPE_GUIDANCE in generation.py.
DECK_TYPE_OPTIONS = {
    "Business / proposal": "business",
    "Technical": "technical",
    "Work status / progress": "status",
    "Casual / D&I / awareness": "casual_di",
    "Presentation / speech": "speech",
}


def _parse_pptx_slides(uploaded_file) -> list[SlideInput]:
    """Read title + body text from an uploaded .pptx. Returns [] if unreadable."""

    try:
        from pptx import Presentation
    except ImportError:
        return []
    try:
        presentation = Presentation(uploaded_file)
    except Exception:
        return []
    slides: list[SlideInput] = []
    for slide in presentation.slides:
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except (AttributeError, ValueError):
            title = ""
        body_parts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if not text or text == title:
                continue
            body_parts.append(text)
        if not title and body_parts:
            title = body_parts.pop(0)
        if title or body_parts:
            slides.append(SlideInput(title or "(untitled)", "\n".join(body_parts)))
    return slides


def _parse_pasted_deck(text: str) -> list[SlideInput]:
    """Split pasted text into slides. Prefer explicit '---' separators; if there
    are none, fall back to blank-line gaps so a whole deck isn't treated as one
    giant slide. First non-empty line of each block is the title."""

    if "---" in text:
        blocks = text.split("---")
    else:
        blocks = re.split(r"\n\s*\n", text)
    slides: list[SlideInput] = []
    for block in blocks:
        lines = [line for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        slides.append(SlideInput(lines[0].strip(), "\n".join(lines[1:]).strip()))
    return slides


def _render_localized_deck(deck, slide_inputs: list[SlideInput]) -> None:
    st.caption(f"Localized with {deck.provider_name}")
    st.subheader("Flow review")
    st.write(deck.flow_review.assessment)
    st.markdown("**Recommended running order:**")
    for suggestion in deck.flow_review.recommended_order:
        st.markdown(
            f"{suggestion.position}. **{suggestion.slide_title}** — {suggestion.rationale}"
        )
    st.divider()
    st.subheader("Slides (copy the Japanese into your deck)")
    for index, localization in enumerate(deck.slides):
        source = slide_inputs[index].title if index < len(slide_inputs) else f"Slide {index + 1}"
        with st.expander(f"Slide {index + 1}: {source}", expanded=index == 0):
            st.markdown("**Japanese title**")
            st.code(localization.japanese_title, language=None)
            st.markdown("**Japanese body**")
            st.code(localization.japanese_body, language=None)
            st.markdown("**Speaker notes (how to deliver it)**")
            st.write(localization.speaker_notes)
            st.markdown("**What changed vs a literal translation, and why**")
            st.write(localization.nuance_notes)


def _clear_present_practice() -> None:
    for key in list(st.session_state):
        if key.startswith("practice_"):
            st.session_state.pop(key, None)


def _practice_reply_input(index: int, options: tuple[str, ...], guided: bool) -> str:
    if guided:
        choice = st.radio(
            "Your delivery",
            options,
            index=None,
            key=f"practice_choice_{index}",
            label_visibility="collapsed",
        )
        return choice or ""
    return st.text_area(
        "Say it in Japanese",
        key=f"practice_text_{index}",
        label_visibility="collapsed",
        height=80,
    ).strip()


def _start_practice(
    generation_service: LessonGenerationService,
    profile: ProfileRecord | None,
    slides: list[SlideInput],
    context: dict[str, str],
    titles: list[str],
) -> bool:
    with st.spinner("Preparing delivery coaching…"):
        try:
            coached = generation_service.coach_deck_delivery(slides, context, profile)
        except GenerationError as error:
            st.error(str(error))
            return False
    _clear_present_practice()
    st.session_state.practice_slides = slides
    st.session_state.practice_titles = titles
    st.session_state.practice_coaching = coached.slides
    st.session_state.practice_provider = coached.provider_name
    st.session_state.practice_context = context
    st.session_state.practice_index = 0
    st.session_state.practice_results = {}
    return True


def render_present_practice(
    generation_service: LessonGenerationService | None,
    profile: ProfileRecord | None,
) -> None:
    """Focused takeover: coach and rehearse delivery, one slide at a time."""

    st.title("Practice presenting")
    if st.button("Leave practice", key="leave_practice"):
        _clear_present_practice()
        st.rerun()
    coaching = st.session_state.get("practice_coaching")
    slides = st.session_state.get("practice_slides")
    if not coaching or not slides or profile is None:
        st.info("This practice session is no longer available.")
        return

    context = st.session_state.get("practice_context", {})
    index = st.session_state.get("practice_index", 0)
    total = len(coaching)
    results = st.session_state.setdefault("practice_results", {})
    provider = st.session_state.get("practice_provider", "")
    st.caption(f"Slide {index + 1} of {total}  ·  coaching by {provider}")

    slide = slides[index]
    coach = coaching[index]
    st.subheader(slide.title)
    if slide.body:
        st.markdown(slide.body.replace("\n", "  \n"))
    st.markdown(f"**Key message:** {coach.key_message}")
    with st.expander("How to deliver this slide", expanded=True):
        st.markdown("**Transition in**")
        st.code(coach.transition_in, language=None)
        st.markdown("**What to say (script)**")
        st.code(coach.spoken_script, language=None)
        st.markdown("**Delivery tips**")
        st.write(coach.delivery_tips)

    st.divider()
    st.markdown("**Now you present this slide:**")
    level = profile.estimated_working_level or profile.declared_level
    guided = level in (
        JapaneseLevel.COMPLETE_BEGINNER,
        JapaneseLevel.N5,
        JapaneseLevel.N4,
        JapaneseLevel.UNSURE,
    )
    reply = _practice_reply_input(index, coach.rehearsal_options, guided)

    if index not in results:
        if st.button("Get feedback", type="primary", key=f"practice_submit_{index}"):
            if not reply:
                st.warning("Say your line first.")
            elif generation_service is None:
                st.error("Feedback is unavailable until model configuration is complete.")
            else:
                with st.spinner("Evaluating your delivery…"):
                    try:
                        judged = generation_service.judge_delivery(
                            coach, slide, reply, context
                        )
                    except GenerationError as error:
                        st.error(str(error))
                    else:
                        results[index] = judged.verdict
                        st.rerun()
        return

    verdict = results[index]
    outcome = evaluate_delivery(verdict)

    def _row(label: str, passed: bool, detail: str) -> None:
        st.markdown(f"{'✅' if passed else '❌'} **{label}** — {detail}")

    _row("Covers the key message", verdict.covers_key_message, verdict.coverage_evidence)
    _row(
        "Comprehensible",
        verdict.comprehensible != Comprehensibility.NO,
        f"{verdict.comprehensible.value}: {verdict.comprehensible_evidence}",
    )
    _row("Right register", verdict.register_appropriate, verdict.register_evidence)
    st.info(f"Tip: {verdict.improvement_tip}")
    if outcome.failed_criteria:
        st.caption("The AI suggests focusing on: " + ", ".join(outcome.failed_criteria) + ".")
    st.caption("AI feedback (beta): it can be wrong, and it is not saved to your progress.")

    left, right = st.columns(2)
    if left.button("Retry this slide", key=f"practice_retry_{index}"):
        results.pop(index, None)
        for stale in (f"practice_choice_{index}", f"practice_text_{index}"):
            st.session_state.pop(stale, None)
        st.rerun()
    if index + 1 < total:
        if right.button("Next slide", type="primary", key=f"practice_next_{index}"):
            st.session_state.practice_index = index + 1
            st.rerun()
    else:
        st.success(f"Deck complete — you rehearsed all {total} slides.")
        st.markdown(
            "**Next: rehearse the whole deck out loud once, then present it for real.**"
        )


def render_presentation(
    generation_service: LessonGenerationService | None = None,
    profile: ProfileRecord | None = None,
) -> None:
    st.title("Presentation")
    st.caption(
        "Turn an English deck into natural Japanese to paste: right register for your audience, "
        "human phrasing, a flow check, and delivery notes."
    )
    if generation_service is None:
        st.info("Presentation localization needs model configuration.")
        return

    if st.session_state.get("practice_coaching"):
        render_present_practice(generation_service, profile)
        return

    audience = st.selectbox(
        "Audience", ("client", "manager", "peer", "subordinate", "executive"), index=0
    )
    formality_col, purpose_col = st.columns(2)
    formality = formality_col.selectbox(
        "Formality", ("formal", "standard business", "casual business"), index=1
    )
    purpose = purpose_col.text_input("Occasion / purpose", value="internal report")
    deck_type_label = st.selectbox("Deck type", tuple(DECK_TYPE_OPTIONS), index=0)
    deck_type = DECK_TYPE_OPTIONS[deck_type_label]
    context = {
        "audience": audience,
        "formality": formality,
        "purpose": purpose,
        "deck_type": deck_type,
    }

    st.markdown("**Your deck**")
    deck_public_ok = _public_content_confirmed("deck_public_ok")
    uploaded = st.file_uploader("Upload a .pptx", type=["pptx"])
    pasted = st.text_area(
        "…or paste slides (separate each slide with a blank line, or a line of ---)",
        height=200,
        placeholder="Q3 Results\nRevenue grew 12%\nAPAC drove growth\n---\nNext Steps\nExpand to two new markets",
    )

    if st.button(
        "Localize to Japanese", type="primary", key="localize_deck", disabled=not deck_public_ok
    ):
        slides: list[SlideInput] = []
        note: str | None = None
        if uploaded is not None:
            slides = _parse_pptx_slides(uploaded)
            if not slides:
                note = "Could not read that .pptx here. Paste the slide text below instead."
        if not slides and pasted.strip():
            slides = _parse_pasted_deck(pasted)
        if not slides:
            st.warning(note or "Upload a .pptx or paste at least one slide.")
        else:
            if len(slides) > PRESENTATION_SLIDE_CAP:
                slides = slides[:PRESENTATION_SLIDE_CAP]
                st.info(f"Localizing the first {PRESENTATION_SLIDE_CAP} slides for this demo.")
            with st.spinner("Localizing your deck…"):
                try:
                    deck = generation_service.localize_deck(slides, context, profile)
                except GenerationError as error:
                    st.error(str(error))
                else:
                    st.session_state.localized_deck = deck
                    st.session_state.localized_slide_inputs = slides

    deck = st.session_state.get("localized_deck")
    if deck is not None:
        _render_localized_deck(deck, st.session_state.get("localized_slide_inputs", []))
        st.divider()
        st.caption("Ready to deliver it? Practice presenting these Japanese slides and get feedback.")
        if st.button("Practice presenting this deck", key="practice_localized"):
            slides = [
                SlideInput(loc.japanese_title, loc.japanese_body) for loc in deck.slides
            ]
            titles = [s.title for s in st.session_state.get("localized_slide_inputs", [])]
            if _start_practice(generation_service, profile, slides, context, titles):
                st.rerun()

    st.divider()
    st.subheader("Practice presenting an existing Japanese deck")
    st.caption("Already have Japanese slides? Upload or paste them and rehearse your delivery.")
    jp_public_ok = _public_content_confirmed("jp_deck_public_ok")
    jp_uploaded = st.file_uploader("Upload a Japanese .pptx", type=["pptx"], key="jp_deck_upload")
    jp_pasted = st.text_area(
        "…or paste Japanese slides (separate each slide with a blank line, or a line of ---)",
        height=150,
        key="jp_deck_paste",
    )
    if st.button("Practice presenting", key="practice_fresh", disabled=not jp_public_ok):
        jp_slides: list[SlideInput] = []
        if jp_uploaded is not None:
            jp_slides = _parse_pptx_slides(jp_uploaded)
        if not jp_slides and jp_pasted.strip():
            jp_slides = _parse_pasted_deck(jp_pasted)
        if not jp_slides:
            st.warning("Upload or paste a Japanese deck first.")
        else:
            if len(jp_slides) > PRESENTATION_SLIDE_CAP:
                jp_slides = jp_slides[:PRESENTATION_SLIDE_CAP]
                st.info(f"Practicing the first {PRESENTATION_SLIDE_CAP} slides.")
            if _start_practice(
                generation_service, profile, jp_slides, context, [s.title for s in jp_slides]
            ):
                st.rerun()


def render_translate() -> None:
    st.title("Translate")
    st.info("Workplace translation will be added in a later increment.")


def _week_label(week) -> str:
    return f"{week:%b} {week.day}"


def _render_over_time(history) -> None:
    """Change over time (decision 15): mastered items per week and recall after a gap."""

    if not history.weeks:
        return
    st.subheader("Over time")
    st.caption(
        "Items mastered: words and grammar points at mastery 0.80 or more at the end of each "
        "week (weeks start Monday, Japan time)."
    )
    mastered_rows = [
        {"Week": _week_label(week.week), "Items mastered": week.mastered} for week in history.weeks
    ]
    top = max(1, max(row["Items mastered"] for row in mastered_rows))
    st.vega_lite_chart(
        mastered_rows,
        {
            "height": 220,
            "mark": {"type": "line", "strokeWidth": 2, "point": {"filled": True, "size": 64}},
            "encoding": {
                "x": {
                    "field": "Week",
                    "type": "ordinal",
                    "sort": None,
                    "axis": {"labelAngle": 0, "title": "Week from"},
                },
                "y": {
                    "field": "Items mastered",
                    "type": "quantitative",
                    "scale": {"domain": [0, top], "nice": True},
                    "axis": {"format": "d", "tickMinStep": 1, "title": "Items mastered"},
                },
                "tooltip": [
                    {"field": "Week", "type": "ordinal", "title": "Week from"},
                    {"field": "Items mastered", "type": "quantitative"},
                ],
            },
        },
        width="stretch",
    )
    recall = [
        {
            "Week": _week_label(week.week),
            "Recall %": round(week.recall_rate * 100),
            "Answers after a gap": week.recall_total,
        }
        for week in history.weeks
        if week.recall_rate is not None
    ]
    if recall:
        st.caption(
            "Recall: share of answers right when the item was last seen 12 hours or more before. "
            "Remembering after a gap is the clearest sign of learning."
        )
        st.vega_lite_chart(
            recall,
            {
                "height": 200,
                "mark": {"type": "bar", "cornerRadiusEnd": 4, "width": {"band": 0.6}},
                "encoding": {
                    "x": {"field": "Week", "type": "ordinal", "sort": None, "axis": {"labelAngle": 0, "title": "Week from"}},
                    "y": {
                        "field": "Recall %",
                        "type": "quantitative",
                        "scale": {"domain": [0, 100]},
                        "axis": {"format": "d", "title": "Recalled (%)"},
                    },
                    "tooltip": [
                        {"field": "Week", "type": "ordinal", "title": "Week from"},
                        {"field": "Recall %", "type": "quantitative", "title": "Recalled (%)"},
                        {"field": "Answers after a gap", "type": "quantitative"},
                    ],
                },
            },
            width="stretch",
        )
    with st.expander("Weekly numbers as a table"):
        st.dataframe(
            [
                {
                    "Week from": week.week.isoformat(),
                    "Items mastered": week.mastered,
                    "Recall answers": week.recall_total,
                    "Recalled": week.recall_correct,
                    "Recall %": (
                        f"{week.recall_rate:.0%}" if week.recall_rate is not None else "-"
                    ),
                }
                for week in history.weeks
            ],
            hide_index=True,
            width="stretch",
        )


def _render_this_week(service: LessonService, user_id: int) -> None:
    """The Home card (decision 15): what moved this week."""

    answers = service.get_graded_answers(user_id)
    if not answers:
        return
    now = service.now()
    week = build_progress_history(answers, now).this_week
    slipping = sum(
        item_status(record, now) is ItemStatus.SLIPPING for record in service.get_progress(user_id)
    )
    st.subheader("This week")
    newly, recalled, answered, slipping_column = st.columns(4)
    newly.metric("Newly mastered", week.newly_mastered)
    recalled.metric(
        "Recalled after a gap",
        f"{week.recall_rate:.0%}" if week.recall_rate is not None else "-",
        help="Right answers on items last seen 12 hours or more before.",
    )
    answered.metric("Answers", week.answers)
    slipping_column.metric("Slipping", slipping)
    if week.recall_total:
        st.caption(f"{week.recall_correct} of {week.recall_total} answers after a gap were right.")


def render_progress(
    service: LessonService | None = None,
    user_id: int | None = None,
    toolkit: LanguageToolkit | None = None,
) -> None:
    st.title("Progress")
    if service is None or user_id is None:
        st.info("Sign in to view progress.")
        return
    progress = service.get_progress(user_id)
    if not progress:
        st.info("Start a lesson to begin tracking your progress.")
        return

    milestones = load_milestones()
    view = build_progress_view(progress, service.now(), toolkit, milestones)
    history = build_progress_history(
        service.get_graded_answers(user_id), service.now(), toolkit, milestones
    )
    columns = st.columns(len(view.status_counts))
    for column, (status, count) in zip(columns, view.status_counts.items()):
        column.metric(status.value.title(), count)
    st.caption(
        "Mastered = recalled correctly on separate days in two question types. "
        "Slipping = answered wrong last time, or a review is well overdue."
    )
    _render_over_time(history)

    if view.milestones:
        st.subheader("Milestones")
        st.caption(
            "Each milestone is a finish line. The word and grammar part is shown here; "
            "reading, writing and presentation checks come in the pilot."
        )
        for item in view.milestones:
            milestone = item.milestone
            st.markdown(f"**{milestone.name}** · {milestone.cefr_hint} · {milestone.summary}")
            st.progress(
                min(1.0, item.share),
                text=(
                    f"{item.mastered} of {item.total} words and grammar points mastered "
                    f"({item.share:.0%}, target {milestone.coverage_target:.0%})"
                    + (
                        f" · {history.milestone_change[milestone.id] * 100:+.1f} points in 4 weeks"
                        if milestone.id in history.milestone_change
                        else ""
                    )
                ),
            )
            with st.expander(f"What {milestone.name} covers"):
                for kind, label in (("read", "Read"), ("write", "Write"), ("speak", "Speak")):
                    tasks = milestone.tasks.get(kind, ())
                    if tasks:
                        st.write(f"{label}: " + "; ".join(tasks))
                st.write(
                    "Word sets: "
                    + ", ".join(f"{part.label} {part.mastered}/{part.total}" for part in item.parts)
                )

    if view.coverage:
        st.subheader("Coverage")
        st.dataframe(
            [
                {
                    "Set": part.label,
                    "Mastered": part.mastered,
                    "Total": part.total,
                    "Share": f"{part.share:.0%}",
                }
                for part in view.coverage
            ],
            hide_index=True,
            width="stretch",
        )

    st.subheader("Your items")
    st.dataframe(
        [
            {
                "Item": row.name,
                "Type": row.category,
                "Status": row.status.value.title(),
                "Mastery": f"{row.mastery:.2f}",
                "Next review": (
                    row.next_review_at.strftime("%Y-%m-%d")
                    if row.next_review_at is not None
                    else "Not scheduled"
                ),
            }
            for row in view.items
        ],
        hide_index=True,
        width="stretch",
    )
    with st.expander("Scheduler details (for reviewers)"):
        st.dataframe(
            [
                {
                    "Item": record.expression or record.item_id,
                    "Exposures": record.exposure_count,
                    "Correct": record.correct_count,
                    "Incorrect": record.incorrect_count,
                    "Outcome": (
                        record.last_outcome.value.title()
                        if record.last_outcome is not None
                        else "No answers"
                    ),
                    "Interval": f"{record.sm2_interval_days} days",
                    "Ease": f"{record.sm2_ease:.2f}",
                    "Successful reviews": record.consecutive_successful_reviews,
                }
                for record in progress
            ],
            hide_index=True,
            width="stretch",
        )
    completions = service.get_completions(user_id)
    if completions:
        latest = completions[0]
        st.caption(
            f"Latest completed topic: {latest.topic_id} · {latest.difficulty} · "
            f"{latest.completed_at.isoformat(sep=' ', timespec='minutes')}"
        )


def _set_suggested_tasks(role_key: str, tasks_key: str) -> None:
    selected_role = st.session_state.get(role_key, "")
    st.session_state[tasks_key] = list(ROLE_TASK_SUGGESTIONS.get(selected_role, ()))


def render_profile_editor(
    service: ProfileService,
    user_id: int,
    profile: ProfileRecord | None,
) -> None:
    onboarding = profile is None
    prefix = f"{'onboarding' if onboarding else 'profile'}_{user_id}"
    role_key = f"{prefix}_role"
    tasks_key = f"{prefix}_tasks"

    if onboarding:
        st.title("Set up your learner profile")
        st.write("Tell us about your work so lessons can match your day-to-day needs.")
    else:
        st.title("Profile")
        st.caption("Update the details used to personalize future lessons.")

    role_options = list(ROLE_TASK_SUGGESTIONS)
    if profile is not None and profile.role not in role_options:
        role_options.insert(0, profile.role)
    selected_role = st.selectbox(
        "Role or title",
        role_options,
        index=role_options.index(profile.role) if profile is not None else None,
        accept_new_options=True,
        key=role_key,
        help="Search common roles or enter your own.",
        on_change=(
            _set_suggested_tasks
            if onboarding
            else None
        ),
        args=(role_key, tasks_key) if onboarding else None,
    )

    if onboarding:
        tasks = []
        tools_domain = None
    else:
        suggested_tasks = list(ROLE_TASK_SUGGESTIONS.get(selected_role, ()))
        selected_tasks = list(profile.tasks)
        task_options = list(dict.fromkeys([*selected_tasks, *suggested_tasks]))
        tasks = st.multiselect(
            "Typical tasks",
            task_options,
            default=None if tasks_key in st.session_state else selected_tasks,
            accept_new_options=True,
            key=tasks_key,
            help="Add the tasks you want lessons to cover.",
        )
        tools_domain = st.text_area(
            "Technologies, tools, or business domain (optional)",
            value=profile.tools_domain or "",
            max_chars=2000,
            key=f"{prefix}_tools_domain",
        )
    level_values = [level.value for level in JapaneseLevel]
    declared_level = st.selectbox(
        "Self-reported Japanese level",
        level_values,
        index=(
            level_values.index(profile.declared_level.value)
            if profile is not None
            else 0
        ),
        key=f"{prefix}_declared_level",
    )
    milestones = load_milestones()
    milestone_names = [milestone.name for milestone in milestones]
    selected_milestone = st.selectbox(
        "Your goal",
        milestone_names,
        index=next(
            (index for index, milestone in enumerate(milestones)
             if profile is not None and milestone.id == profile.milestone_id),
            0,
        ),
        key=f"{prefix}_milestone",
    )
    milestone_id = milestones[milestone_names.index(selected_milestone)].id
    romaji_preference = st.toggle(
        "Show romaji support",
        value=profile.romaji_preference if profile is not None else False,
        key=f"{prefix}_romaji",
        help="Optional reading support, especially for complete beginners.",
    )

    if profile is None or profile.estimated_working_level is None:
        st.caption("Estimated working level: Not set")
    else:
        st.caption(f"Estimated working level: {profile.estimated_working_level.value}")
    st.button(
        "Take optional placement",
        disabled=True,
        help="Optional placement will be available in a later increment.",
        key=f"{prefix}_placement",
    )

    submit_label = "Save profile" if onboarding else "Save changes"
    if st.button(submit_label, type="primary", key=f"{prefix}_submit"):
        try:
            if onboarding:
                service.create_profile(
                    user_id,
                    role=selected_role or "",
                    tasks=tasks,
                    tools_domain=tools_domain,
                    declared_level=declared_level,
                    milestone_id=milestone_id,
                    romaji_preference=romaji_preference,
                )
            else:
                service.update_profile(
                    user_id,
                    role=selected_role or "",
                    tasks=tasks,
                    tools_domain=tools_domain,
                    declared_level=declared_level,
                    milestone_id=milestone_id,
                    romaji_preference=romaji_preference,
                )
        except ProfileValidationError as error:
            st.error(str(error))
        else:
            if onboarding:
                st.rerun()
            st.success("Profile saved.")


def render_profile(
    service: ProfileService | None = None,
    user_id: int | None = None,
    profile: ProfileRecord | None = None,
    term_service: PersonalTermService | None = None,
) -> None:
    if service is None or user_id is None:
        st.title("Profile")
        st.info("Sign in to view your learner profile.")
        return
    render_profile_editor(service, user_id, profile)
    if profile is not None and term_service is not None:
        render_personal_terms(term_service, user_id)


def render_personal_terms(service: PersonalTermService, user_id: int) -> None:
    st.divider()
    st.subheader("Your words and grammar")
    st.caption("Add public Japanese terms you want to practise. Readings come from the reference dictionary.")
    kind = st.selectbox("Add", ("word", "grammar"), key="personal_terms_kind")
    importance = st.selectbox("Importance", ("A", "S", "B"), key="personal_terms_importance")
    pasted = st.text_area(
        "One term per line, optionally followed by a comma and meaning",
        key="personal_terms_paste",
        max_chars=50000,
    )
    if st.button("Preview terms", key="preview_personal_terms"):
        try:
            preview = service.preview(user_id, pasted, kind=kind, importance=importance)
        except TermValidationError as error:
            st.error(str(error))
        else:
            st.session_state.personal_terms_preview = (pasted, kind, importance, preview)
    current = st.session_state.get("personal_terms_preview")
    if current is not None and current[:3] == (pasted, kind, importance):
        preview = current[3]
        st.dataframe(
            [{"Term": row.expression, "Status": row.status,
              "Reading": row.reading or "—", "Importance": row.importance}
             for row in preview],
            hide_index=True,
        )
        confirmed = st.checkbox(
            "These terms are public and contain no client, project or internal names",
            key="personal_terms_public",
        )
        if st.button("Save terms", key="save_personal_terms", disabled=not confirmed):
            try:
                count = service.save(user_id, preview, public_confirmed=confirmed)
            except TermValidationError as error:
                st.error(str(error))
            else:
                st.session_state.pop("personal_terms_preview", None)
                st.success(f"Saved {count} term{'s' if count != 1 else ''}.")
    terms = service.list_terms(user_id)
    if terms:
        choices = {f"{term.expression} ({term.kind}) · {term.id}": term for term in terms}
        label = st.selectbox("Manage a term", list(choices), key="manage_personal_term")
        term = choices[label]
        grade = st.selectbox(
            "Importance for this term", ("S", "A", "B"),
            index=("S", "A", "B").index(term.importance), key="manage_term_importance",
        )
        hidden = st.checkbox("Hide from lesson targets", value=term.status == "hidden",
                             key="manage_term_hidden")
        if st.button("Update term", key="update_personal_term"):
            service.set_importance(user_id, term.id, grade)
            service.set_hidden(user_id, term.id, hidden)
            st.rerun()


PAGE_RENDERERS: dict[str, PageRenderer] = {
    "Home": render_home,
    "Learn": render_learn,
    "Presentation": render_presentation,
    "Translate": render_translate,
    "Progress": render_progress,
    "Profile": render_profile,
}


def render_health_panel(settings: Settings, engine: Engine) -> None:
    """Show operational readiness without revealing configuration values."""

    health = check_database(engine)

    with st.sidebar.expander("Startup health", expanded=True):
        if health.ready:
            st.success("Database: ready")
            st.caption("SQLite foreign keys: enabled")
        else:
            st.error(f"Database: {health.message.lower()}")

        if settings.llm_mode == "mock":
            st.info("Model features: mock mode (offline fixtures, no model calls)")
        elif settings.model_configured:
            st.success("Model features: configured")
        else:
            st.warning("Model features unavailable: configuration is missing")
        st.caption("Secret values are never displayed.")


def render_authentication(service: AuthenticationService) -> bool:
    st.title("Demo authentication")
    st.caption("Local username and password access for this laptop demo.")
    authenticated = False

    sign_in_tab, register_tab = st.tabs(("Sign in", "Register"))
    with sign_in_tab:
        with st.form("sign_in_form"):
            username = st.text_input("Username", key="sign_in_username")
            password = st.text_input("Password", type="password", key="sign_in_password")
            sign_in = st.form_submit_button("Sign in", type="primary")
        if sign_in:
            try:
                user = service.authenticate(username, password)
            except AuthenticationError as error:
                st.error(str(error))
            else:
                st.session_state.authenticated_user_id = user.id
                st.session_state.authenticated_username = user.username
                authenticated = True

    with register_tab:
        with st.form("registration_form"):
            username = st.text_input("Username", key="registration_username")
            password = st.text_input(
                "Password",
                type="password",
                key="registration_password",
                help="Use 12-128 characters.",
            )
            register = st.form_submit_button("Create account", type="primary")
        if register:
            try:
                user = service.register(username, password)
            except RegistrationError as error:
                st.error(str(error))
            else:
                st.session_state.authenticated_user_id = user.id
                st.session_state.authenticated_username = user.username
                authenticated = True

    return authenticated


def clear_authenticated_session() -> None:
    _clear_active_lesson()
    _clear_active_review()
    st.session_state.pop("last_lesson_completion", None)
    st.session_state.pop("last_review_summary", None)
    st.session_state.pop("authenticated_user_id", None)
    st.session_state.pop("authenticated_username", None)
    st.session_state.pop("personal_terms_preview", None)
    for key in ("today_route", "today_scenario", "lesson_scenario",
                "pending_navigation"):
        st.session_state.pop(key, None)
    for key in list(st.session_state):
        if key.startswith(("personal_terms_", "manage_personal_term", "manage_term_",
                           "personal_focus_")):
            st.session_state.pop(key, None)


def main() -> None:
    settings = get_settings()
    st.set_page_config(page_title=settings.app_name, page_icon="🎌", layout="wide")
    st.markdown(JAPANESE_STYLE, unsafe_allow_html=True)
    project_root = Path(__file__).resolve().parents[2]
    engine = create_database_engine(settings, project_root)
    st.sidebar.title(settings.app_name)
    render_health_panel(settings, engine)

    service = AuthenticationService(engine)
    if "authenticated_user_id" not in st.session_state:
        authenticated = render_authentication(service)
        engine.dispose()
        if authenticated:
            st.rerun()
        return

    st.sidebar.caption(
        f"Signed in as {st.session_state.get('authenticated_username', 'local user')}"
    )
    if st.sidebar.button("Sign out"):
        clear_authenticated_session()
        engine.dispose()
        st.rerun()

    user_id = int(st.session_state.authenticated_user_id)
    profile_service = ProfileService(engine)
    toolkit = SudachiToolkit.shared()
    lesson_service = LessonService(engine, toolkit=toolkit)
    generation_service = None
    if settings.llm_mode == "mock":
        generation_service = _mock_generation_service(settings.mock_fixture_dir)
    elif settings.model_configured:
        api_key = settings.model_api_key or settings.apim_subscription_key
        if (
            settings.model_base_url is not None
            and api_key is not None
            and settings.primary_model is not None
            and settings.fallback_model is not None
        ):
            generation_service = _generation_service(
                settings.model_base_url,
                api_key.get_secret_value(),
                settings.primary_model,
                settings.fallback_model,
                settings.model_timeout_seconds,
                settings.primary_model_timeout_seconds,
                settings.hedge_after_seconds,
                settings.parallel_explanations,
                settings.explanation_workers,
                # Route the strict-JSON quiz phase to the structured (fallback) model
                # first, then the primary. Overridable via JLT_QUIZ_* settings.
                settings.quiz_primary_model or settings.fallback_model,
                settings.quiz_fallback_model or settings.primary_model,
                settings.llm_mode,
                settings.llm_recording_path,
                settings.model_provider,
                settings.azure_api_version,
                (settings.apim_subscription_key.get_secret_value()
                 if settings.apim_subscription_key else None),
                settings.apim_subscription_key_header,
            )
    profile = profile_service.get_profile(user_id)
    if profile is None:
        render_profile_editor(profile_service, user_id, None)
        engine.dispose()
        return

    review_service = ReviewService(engine) if settings.reviewer_mode else None
    pages = [*PAGE_RENDERERS.keys(), *(["Accuracy"] if review_service else [])]
    pending_navigation = st.session_state.pop("pending_navigation", None)
    if pending_navigation in pages:
        st.session_state.navigation = pending_navigation
    selected_page = st.sidebar.radio("Navigation", pages, key="navigation")
    if selected_page == "Home":
        render_home(profile, lesson_service, user_id,
                    TodayService(profile_service, lesson_service, toolkit))
    elif selected_page == "Learn":
        render_learn(lesson_service, user_id, profile, generation_service,
                     review_service, PersonalTermService(engine, toolkit),
                     SituationWordService(toolkit) if settings.llm_mode != "mock" else None,
                     DomainWordService(toolkit) if settings.llm_mode != "mock" else None,
                     settings.llm_mode,
                     TodayService(profile_service, lesson_service, toolkit))
    elif selected_page == "Presentation":
        render_presentation(generation_service, profile)
    elif selected_page == "Progress":
        render_progress(lesson_service, user_id, toolkit)
    elif selected_page == "Profile":
        render_profile(profile_service, user_id, profile,
                       PersonalTermService(engine, toolkit))
    elif selected_page == "Accuracy" and review_service is not None:
        render_accuracy(review_service)
    else:
        PAGE_RENDERERS[selected_page]()
    engine.dispose()


if __name__ == "__main__":
    main()
